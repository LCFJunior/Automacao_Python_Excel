# Import

import re
import pandas as pd
from docx import Document
from openpyxl import load_workbook


# Função para extrair informações dos sócios

def extract_partner_info(text, gender):
    if gender == 'male':
        pattern = re.compile(
            r'\d+\.\s+([\w\s]+),\s+portador\s+do\s+CPF\s+\d{3}\.\d{3}\.\d{3}-\d{2},\s+residente\s+em\s+[\w\s,]+,\s+doravante\s+denominado\s+"Sócio\s+\d+",\s+detentor\s+de\s+(\d+)\s+cotas?\.'
        )
    elif gender == 'female':
        pattern = re.compile(
            r'\d+\.\s+([\w\s]+),\s+portadora\s+do\s+CPF\s+\d{3}\.\d{3}\.\d{3}-\d{2},\s+residente\s+em\s+[\w\s,]+,\s+doravante\s+denominada\s+"Sócio\s+\d+",\s+detentora\s+de\s+(\d+)\s+cotas?\.'
        )
    matches = pattern.findall(text)
    return [(match[0].strip(), int(match[1])) for match in matches]


# Restante do processo

doc = Document(r'C:\Users\Junior\Desktop\AAWZ_Partners\Partnership.docx')

full_text = "\n".join([para.text for para in doc.paragraphs])
male_partners = extract_partner_info(full_text, 'male')
female_partners = extract_partner_info(full_text, 'female')
all_partners = male_partners + female_partners

df = pd.DataFrame(all_partners, columns=['Nome', 'Cotas'])


# Salvar

file_path = r'C:\Users\Junior\Desktop\AAWZ_Partners\Partnership_V2.xlsx'
df.to_excel(file_path, index=False)


# Formatar

wb = load_workbook(file_path)
ws = wb.active
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 20
wb.save(file_path)
